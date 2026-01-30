"""
Views e utilitários para integração com OnlyOffice Document Server
para Norma Jurídica
"""
import hashlib
import json
import logging
import time
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.contrib import messages

from sapl.norma.models import NormaJuridica

logger = logging.getLogger(__name__)


def generate_file_key(norma_id, user_id):
    """
    Gera uma chave única para o documento no OnlyOffice
    A chave muda a cada edição para forçar o OnlyOffice a recarregar
    """
    timestamp = str(int(time.time()))
    string_to_hash = f"norma_{norma_id}_user_{user_id}_{timestamp}"
    return hashlib.md5(string_to_hash.encode()).hexdigest()


@login_required
@require_http_methods(["GET"])
def norma_onlyoffice_config(request, pk):
    """
    Retorna a configuração JSON para inicializar o editor OnlyOffice
    """
    norma = get_object_or_404(NormaJuridica, pk=pk)

    # Verifica permissão de edição
    can_edit = request.user.has_perm('norma.change_normajuridica')

    # URLs para o OnlyOffice acessar (dentro da rede Docker)
    download_url = request.build_absolute_uri(
        reverse('sapl.norma:norma_onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = request.build_absolute_uri(
        reverse('sapl.norma:norma_onlyoffice_callback', kwargs={'pk': pk})
    )

    # Substituir localhost/host externo pelo nome do container na rede Docker
    host = request.get_host()
    download_url = download_url.replace(f'http://{host}', 'http://sapl-dev:8000')
    download_url = download_url.replace(f'https://{host}', 'http://sapl-dev:8000')
    callback_url = callback_url.replace(f'http://{host}', 'http://sapl-dev:8000')
    callback_url = callback_url.replace(f'https://{host}', 'http://sapl-dev:8000')

    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key(norma.pk, request.user.pk),
        "title": f"Norma_{norma.pk}.docx",
        "url": download_url,
    }

    # Configuração do editor
    editor_config = {
        "mode": "edit" if can_edit else "view",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@require_http_methods(["GET"])
def norma_onlyoffice_download(request, pk):
    """
    Endpoint para o OnlyOffice baixar o documento
    Se não existe arquivo, retorna um documento em branco
    """
    norma = get_object_or_404(NormaJuridica, pk=pk)

    # Se já tem arquivo, retorna ele
    if norma.texto_integral:
        try:
            with open(norma.texto_integral.path, 'rb') as f:
                content = f.read()
            response = HttpResponse(content, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename="Norma_{pk}.docx"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo: {e}")

    # Tenta criar documento usando template
    try:
        from sapl.utils_template import criar_documento_com_template, criar_documento_em_branco

        # Busca tipo específico da norma (TipoNormaJuridica)
        tipo_especifico = norma.tipo if norma.tipo else None

        file_stream = criar_documento_com_template(
            'norma',
            tipo_especifico,
            {
                'titulo': f'{norma.tipo} {norma.numero}/{norma.ano}',
                'descricao': norma.ementa,
                'tipo_display': 'Norma Jurídica'
            }
        )

        # Se não encontrou template, cria documento em branco
        if not file_stream:
            file_stream = criar_documento_em_branco(
                f'{norma.tipo} {norma.numero}/{norma.ano}',
                norma.ementa,
                'Norma Jurídica'
            )

        if file_stream:
            response = HttpResponse(
                file_stream.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="Norma_{pk}.docx"'
            return response

    except Exception as e:
        logger.error(f"Erro ao criar documento: {e}")

    # Fallback: Se python-docx não está instalado ou houve erro
    try:
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_heading(f'{norma.tipo} {norma.numero}/{norma.ano}', 0)
        doc.add_paragraph(f'Ementa: {norma.ementa}')
        doc.add_paragraph('')
        doc.add_paragraph('Digite o texto da norma abaixo:')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="Norma_{pk}.docx"'
        return response

    except ImportError:
        logger.error("python-docx não está instalado")
        return HttpResponse("Erro: python-docx não instalado", status=500)


@csrf_exempt
@require_http_methods(["POST"])
def norma_onlyoffice_callback(request, pk):
    """
    Callback chamado pelo OnlyOffice quando o documento é salvo
    """
    try:
        body = json.loads(request.body.decode('utf-8'))
        status = body.get('status')
        download_url = body.get('url')

        logger.info(f"OnlyOffice callback para norma jurídica {pk}: status={status}, url={download_url}")

        # Status 2 ou 6 significa que o documento foi salvo
        if status in [2, 6] and download_url:
            # Substitui localhost:8001 por onlyoffice:80 para acesso interno Docker
            if 'localhost:8001' in download_url:
                download_url = download_url.replace('localhost:8001', 'onlyoffice:80')

            norma = get_object_or_404(NormaJuridica, pk=pk)

            # Baixa o documento do OnlyOffice
            import requests
            try:
                response = requests.get(download_url, timeout=30)
            except Exception as e:
                logger.error(f"Erro ao fazer requisição de download: {e}")
                return JsonResponse({"error": 1})

            if response.status_code == 200:
                from django.core.files.base import ContentFile

                filename = f"norma_{pk}_{int(time.time())}.docx"

                # Remove arquivo antigo se existir
                if norma.texto_integral:
                    norma.texto_integral.delete(save=False)

                try:
                    norma.texto_integral.save(
                        filename,
                        ContentFile(response.content),
                        save=True
                    )
                    logger.info(f"Documento salvo com sucesso: {filename}")
                    return JsonResponse({"error": 0})
                except Exception as e:
                    logger.error(f"Erro ao salvar arquivo: {e}")
                    return JsonResponse({"error": 1})
            else:
                logger.error(f"Erro ao baixar documento: status={response.status_code}")
                return JsonResponse({"error": 1})

        return JsonResponse({"error": 0})

    except Exception as e:
        logger.error(f"Erro no callback OnlyOffice: {e}")
        return JsonResponse({"error": 1})


@login_required
def norma_onlyoffice_editor(request, pk):
    """
    Renderiza a página com o editor OnlyOffice integrado
    """
    norma = get_object_or_404(NormaJuridica, pk=pk)

    # Verifica se o usuário tem permissão
    if not request.user.has_perm('norma.change_normajuridica'):
        messages.error(request, 'Você não tem permissão para editar esta norma.')
        return redirect('sapl.norma:normajuridica_detail', pk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = settings.ONLYOFFICE_URL
    if 'onlyoffice:' in onlyoffice_url or 'onlyoffice/' in onlyoffice_url:
        protocol = 'https' if request.is_secure() else 'http'
        host = request.get_host().split(':')[0]
        onlyoffice_url = f"{protocol}://{host}:8001"

    context = {
        'documento': norma,
        'documento_tipo': 'Norma Jurídica',
        'documento_titulo': f'{norma.tipo} {norma.numero}/{norma.ano}',
        'documento_descricao': norma.ementa,
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.norma:norma_onlyoffice_config', kwargs={'pk': pk}),
        'voltar_url': reverse('sapl.norma:normajuridica_detail', kwargs={'pk': pk}),
    }

    return render(request, 'onlyoffice/onlyoffice_editor.html', context)
