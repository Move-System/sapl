"""
Views e utilitários para integração com OnlyOffice Document Server
para edição de Templates de Documentos
"""
import hashlib
import json
import logging
import time
from django.conf import settings
from django.contrib.auth.decorators import login_required, permission_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.contrib import messages

from sapl.base.models import DocumentTemplate

logger = logging.getLogger(__name__)


def generate_file_key(template_id, user_id):
    """
    Gera uma chave única para o documento no OnlyOffice
    A chave muda a cada edição para forçar o OnlyOffice a recarregar
    """
    timestamp = str(int(time.time()))
    string_to_hash = f"template_{template_id}_user_{user_id}_{timestamp}"
    return hashlib.md5(string_to_hash.encode()).hexdigest()


@login_required
@permission_required('base.change_documenttemplate', raise_exception=True)
@require_http_methods(["GET"])
def template_onlyoffice_config(request, pk):
    """
    Retorna a configuração JSON para inicializar o editor OnlyOffice
    """
    template = get_object_or_404(DocumentTemplate, pk=pk)

    # URLs para o OnlyOffice acessar (dentro da rede Docker)
    download_url = request.build_absolute_uri(
        reverse('sapl.base:template_onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = request.build_absolute_uri(
        reverse('sapl.base:template_onlyoffice_callback', kwargs={'pk': pk})
    )


    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key(template.pk, request.user.pk),
        "title": f"Template_{template.pk}.docx",
        "url": download_url,
    }

    # Configuração do editor
    editor_config = {
        "mode": "edit",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@require_http_methods(["GET"])
def template_onlyoffice_download(request, pk):
    """
    Endpoint para o OnlyOffice baixar o documento do template
    Se não existe arquivo, retorna um documento em branco
    """
    template = get_object_or_404(DocumentTemplate, pk=pk)

    # Se já tem arquivo, retorna ele
    if template.arquivo:
        try:
            with open(template.arquivo.path, 'rb') as f:
                content = f.read()
            response = HttpResponse(
                content,
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="Template_{pk}.docx"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo do template: {e}")

    # Se não tem arquivo, cria um documento em branco usando python-docx
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from io import BytesIO

        doc = Document()

        # Adiciona cabeçalho padrão
        section = doc.sections[0]
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = f"[CABEÇALHO - {template.nome}]"
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Adiciona rodapé padrão
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = "[RODAPÉ]"
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Adiciona conteúdo inicial
        doc.add_heading(f'Template: {template.nome}', 0)
        doc.add_paragraph(f'Tipo: {template.get_tipo_conteudo_display()}')
        if template.descricao:
            doc.add_paragraph(f'Descrição: {template.descricao}')
        doc.add_paragraph('')
        doc.add_paragraph('Edite este documento para configurar o template.')
        doc.add_paragraph('O cabeçalho e rodapé definidos aqui serão aplicados aos novos documentos.')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="Template_{pk}.docx"'
        return response

    except ImportError:
        logger.error("python-docx não está instalado")
        return HttpResponse("Erro: python-docx não instalado", status=500)


@csrf_exempt
@require_http_methods(["POST"])
def template_onlyoffice_callback(request, pk):
    """
    Callback chamado pelo OnlyOffice quando o documento é salvo
    """
    try:
        body = json.loads(request.body.decode('utf-8'))
        status = body.get('status')
        download_url = body.get('url')

        logger.info(f"OnlyOffice callback para template {pk}: status={status}, url={download_url}")

        # Status 2 ou 6 significa que o documento foi salvo
        if status in [2, 6] and download_url:

            template = get_object_or_404(DocumentTemplate, pk=pk)

            # Baixa o documento do OnlyOffice
            import requests
            try:
                response = requests.get(download_url, timeout=30)
            except Exception as e:
                logger.error(f"Erro ao fazer requisição de download: {e}")
                return JsonResponse({"error": 1})

            if response.status_code == 200:
                from django.core.files.base import ContentFile

                filename = f"template_{pk}_{int(time.time())}.docx"

                # Remove arquivo antigo se existir
                if template.arquivo:
                    template.arquivo.delete(save=False)

                try:
                    template.arquivo.save(
                        filename,
                        ContentFile(response.content),
                        save=True
                    )
                    logger.info(f"Template salvo com sucesso: {filename}")
                    return JsonResponse({"error": 0})
                except Exception as e:
                    logger.error(f"Erro ao salvar arquivo: {e}")
                    return JsonResponse({"error": 1})
            else:
                logger.error(f"Erro ao baixar documento: status={response.status_code}")
                return JsonResponse({"error": 1})

        return JsonResponse({"error": 0})

    except Exception as e:
        logger.error(f"Erro no callback OnlyOffice: {e}")
        return JsonResponse({"error": 1})


@login_required
@permission_required('base.change_documenttemplate', raise_exception=True)
def template_onlyoffice_editor(request, pk):
    """
    Renderiza a página com o editor OnlyOffice integrado para editar template
    """
    template = get_object_or_404(DocumentTemplate, pk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = settings.ONLYOFFICE_URL

    context = {
        'documento': template,
        'documento_tipo': 'Template de Documento',
        'documento_titulo': template.nome,
        'documento_descricao': template.descricao or f'Template para {template.get_tipo_conteudo_display()}',
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.base:template_onlyoffice_config', kwargs={'pk': pk}),
        'voltar_url': reverse('sapl.base:documenttemplate_detail', kwargs={'pk': pk}),
    }

    return render(request, 'onlyoffice/onlyoffice_editor.html', context)
