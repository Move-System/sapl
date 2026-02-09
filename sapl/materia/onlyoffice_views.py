"""
Views e utilitários para integração com OnlyOffice Document Server
"""
import hashlib
import json
import logging
import os
import time
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods

from sapl.materia.models import Proposicao
from sapl.utils import get_base_url, build_onlyoffice_url, get_onlyoffice_browser_url

logger = logging.getLogger(__name__)


def generate_file_key(proposicao_id, user_id):
    """
    Gera uma chave única para o documento no OnlyOffice
    A chave muda a cada edição para forçar o OnlyOffice a recarregar
    """
    timestamp = str(int(time.time()))
    string_to_hash = f"proposicao_{proposicao_id}_user_{user_id}_{timestamp}"
    return hashlib.md5(string_to_hash.encode()).hexdigest()


@login_required
@require_http_methods(["GET"])
def onlyoffice_config(request, pk):
    """
    Retorna a configuração JSON para inicializar o editor OnlyOffice
    """
    proposicao = get_object_or_404(Proposicao, pk=pk)

    # Verifica se o usuário tem permissão para editar
    can_edit = (
        not proposicao.data_envio and
        proposicao.autor.operadores.filter(id=request.user.id).exists()
    )

    # URLs para o OnlyOffice acessar (dentro da rede Docker)
    # Usa SAPL_INTERNAL_URL quando configurada para comunicação container-a-container
    download_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:onlyoffice_callback', kwargs={'pk': pk})
    )

    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key(proposicao.pk, request.user.pk),
        "title": f"Proposicao_{proposicao.pk}.docx",
        "url": download_url,
    }

    # Configuração do editor
    editor_config = {
        "mode": "edit" if can_edit else "view",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@require_http_methods(["GET"])
def onlyoffice_download(request, pk):
    """
    Endpoint para o OnlyOffice baixar o documento
    Se não existe arquivo, retorna um documento em branco
    NOTA: Sem @login_required pois o OnlyOffice não tem sessão do Django
    """
    proposicao = get_object_or_404(Proposicao, pk=pk)

    # Verificação de permissão apenas se houver usuário autenticado
    # Permite: operadores do autor OU usuários com permissão de protocolo
    if request.user.is_authenticated:
        is_operator = proposicao.autor.operadores.filter(id=request.user.id).exists()
        has_protocol_perm = request.user.has_perm('protocoloadm.add_documentoadministrativo')
        if not is_operator and not has_protocol_perm:
            return HttpResponse("Sem permissão", status=403)

    # Se já tem arquivo, retorna ele
    if proposicao.texto_original:
        try:
            with open(proposicao.texto_original.path, 'rb') as f:
                content = f.read()
            response = HttpResponse(content, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename="Proposicao_{pk}.docx"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo: {e}")

    # Tenta criar documento usando template
    try:
        from sapl.utils_template import criar_documento_com_template, criar_documento_em_branco

        # Busca tipo específico da proposição (TipoProposicao)
        tipo_especifico = proposicao.tipo if proposicao.tipo else None

        file_stream = criar_documento_com_template(
            'proposicao',
            tipo_especifico,
            {
                'titulo': f'Proposição {proposicao.tipo}',
                'descricao': proposicao.descricao,
                'tipo_display': 'Proposição'
            }
        )

        # Se não encontrou template, cria documento em branco
        if not file_stream:
            file_stream = criar_documento_em_branco(
                f'Proposição {proposicao.tipo}',
                proposicao.descricao,
                'Proposição'
            )

        if file_stream:
            response = HttpResponse(
                file_stream.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="Proposicao_{pk}.docx"'
            return response

    except Exception as e:
        logger.error(f"Erro ao criar documento: {e}")

    # Fallback: Se python-docx não está instalado ou houve erro
    try:
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_heading(f'Proposição {proposicao.tipo}', 0)
        doc.add_paragraph(f'Ementa: {proposicao.descricao}')
        doc.add_paragraph('')
        doc.add_paragraph('Digite o texto da proposição abaixo:')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="Proposicao_{pk}.docx"'
        return response

    except ImportError:
        logger.error("python-docx não está instalado")
        return HttpResponse("Erro: python-docx não instalado", status=500)


@csrf_exempt
@require_http_methods(["POST"])
def onlyoffice_callback(request, pk):
    """
    Callback chamado pelo OnlyOffice quando o documento é salvo

    Status codes:
    0 - Nenhum documento com a chave identificada foi encontrado
    1 - Documento está sendo editado
    2 - Documento está pronto para salvar
    3 - Ocorreu erro ao salvar o documento
    4 - Documento está fechado sem alterações
    6 - Documento está sendo editado, mas a versão atual do documento foi salva
    7 - Ocorreu um erro de salvamento forçado ou salvamento automático
    """
    try:
        body = json.loads(request.body.decode('utf-8'))
        status = body.get('status')
        download_url = body.get('url')

        logger.info(f"OnlyOffice callback para proposição {pk}: status={status}, url={download_url}, body={body}")

        # Status 2 ou 6 significa que o documento foi salvo
        if status in [2, 6] and download_url:
            logger.info(f"URL original recebida: {download_url}")

            # Substitui o host da URL pelo ONLYOFFICE_URL configurado
            # O OnlyOffice retorna URLs com seu próprio hostname que pode não ser acessível
            import re
            from urllib.parse import urlparse
            onlyoffice_parsed = urlparse(settings.ONLYOFFICE_URL)
            onlyoffice_base = f"{onlyoffice_parsed.scheme}://{onlyoffice_parsed.netloc}"
            download_url = re.sub(
                r'https?://[^/]+',
                onlyoffice_base,
                download_url
            )
            logger.info(f"URL substituída para: {download_url}")

            logger.info(f"Iniciando download do documento de: {download_url}")
            proposicao = get_object_or_404(Proposicao, pk=pk)

            # Baixa o documento do OnlyOffice
            import requests
            try:
                response = requests.get(download_url, timeout=30)
                logger.info(f"Download response: status={response.status_code}, size={len(response.content)}")
            except Exception as e:
                logger.error(f"Erro ao fazer requisição de download: {e}")
                return JsonResponse({"error": 1})

            if response.status_code == 200:
                # Salva o arquivo
                from django.core.files.base import ContentFile

                filename = f"proposicao_{pk}_{int(time.time())}.docx"
                logger.info(f"Salvando arquivo: {filename}")

                # Remove arquivo antigo se existir
                if proposicao.texto_original:
                    old_file = proposicao.texto_original.name
                    proposicao.texto_original.delete(save=False)
                    logger.info(f"Arquivo antigo removido: {old_file}")

                try:
                    proposicao.texto_original.save(
                        filename,
                        ContentFile(response.content),
                        save=False
                    )

                    # Regenera o hash_code se a proposição está em confirmação
                    if proposicao.data_envio and not proposicao.data_recebimento:
                        from sapl.utils import gerar_hash_arquivo
                        proposicao.hash_code = gerar_hash_arquivo(
                            proposicao.texto_original.path, str(proposicao.pk))
                        logger.info(f"Hash code atualizado: {proposicao.hash_code}")

                    proposicao.save()
                    logger.info(f"Documento salvo com sucesso: {filename}")
                    return JsonResponse({"error": 0})
                except Exception as e:
                    logger.error(f"Erro ao salvar arquivo: {e}")
                    return JsonResponse({"error": 1})
            else:
                logger.error(f"Erro ao baixar documento: status={response.status_code}")
                return JsonResponse({"error": 1})

        # Para outros status, apenas retorna sucesso
        return JsonResponse({"error": 0})

    except Exception as e:
        logger.error(f"Erro no callback OnlyOffice: {e}")
        return JsonResponse({"error": 1})


@login_required
def onlyoffice_editor(request, pk):
    """
    Renderiza a página com o editor OnlyOffice integrado
    """
    from django.shortcuts import render

    proposicao = get_object_or_404(Proposicao, pk=pk)

    # Verifica se o usuário tem permissão
    if not proposicao.autor.operadores.filter(id=request.user.id).exists():
        from django.contrib import messages
        from django.shortcuts import redirect
        messages.error(request, 'Você não tem permissão para editar esta proposição.')
        return redirect('sapl.materia:proposicao_detail', pk=pk)

    # Verifica se já foi enviada
    if proposicao.data_envio:
        from django.contrib import messages
        from django.shortcuts import redirect
        messages.warning(request, 'Esta proposição já foi enviada e não pode mais ser editada.')
        return redirect('sapl.materia:proposicao_detail', pk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = get_onlyoffice_browser_url(request)

    context = {
        'proposicao': proposicao,
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.materia:onlyoffice_config', kwargs={'pk': pk}),
    }

    return render(request, 'materia/onlyoffice_editor.html', context)


@login_required
@require_http_methods(["GET"])
def onlyoffice_confirmar_config(request, pk):
    """
    Retorna a configuração JSON para o editor OnlyOffice na página de confirmação
    """
    proposicao = get_object_or_404(Proposicao, pk=pk)

    # Verifica se o usuário tem permissão de protocolo
    if not request.user.has_perm('protocoloadm.add_documentoadministrativo'):
        return JsonResponse({"error": "Sem permissão"}, status=403)

    # Verifica se a proposição está em estado de confirmação
    # (enviada mas não recebida)
    if not proposicao.data_envio or proposicao.data_recebimento:
        return JsonResponse({"error": "Proposição não está em confirmação"}, status=400)

    # URLs para o OnlyOffice acessar o documento
    download_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:onlyoffice_callback', kwargs={'pk': pk})
    )

    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key(proposicao.pk, request.user.pk),
        "title": f"Proposicao_{proposicao.pk}.docx",
        "url": download_url,
        "permissions": {
            "edit": True,
            "download": True,
            "print": True,
            "review": True,
            "comment": True,
        },
    }

    # Configuração do editor - permite edição na confirmação
    editor_config = {
        "mode": "edit",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@login_required
def onlyoffice_confirmar_editor(request, pk, hash):
    """
    Renderiza a página com o editor OnlyOffice para proposição em confirmação
    """
    from django.shortcuts import render, redirect
    from django.contrib import messages
    from sapl.materia.forms import SEPARADOR_HASH_PROPOSICAO
    from sapl.utils import gerar_hash_arquivo

    proposicao = get_object_or_404(Proposicao, pk=pk)

    # Verifica se o usuário tem permissão de protocolo
    if not request.user.has_perm('protocoloadm.add_documentoadministrativo'):
        messages.error(request, 'Você não tem permissão para editar esta proposição.')
        return redirect('sapl.materia:proposicao-confirmar', hash=hash, pk=pk)

    # Verifica se a proposição está em estado de confirmação
    if not proposicao.data_envio or proposicao.data_recebimento:
        messages.warning(request, 'Esta proposição não está em estado de confirmação.')
        return redirect('sapl.materia:proposicao-confirmar', hash=hash, pk=pk)

    # Verifica o hash (mesma lógica de ConfirmarProposicao)
    if proposicao.texto_articulado.exists():
        ta = proposicao.texto_articulado.first()
        hasher = 'P' + ta.hash() + SEPARADOR_HASH_PROPOSICAO + str(proposicao.pk)
    else:
        hasher = gerar_hash_arquivo(
            proposicao.texto_original.path,
            str(proposicao.pk)) if proposicao.texto_original else None

    expected_hash = 'P%s%s%s' % (hash, SEPARADOR_HASH_PROPOSICAO, proposicao.pk)
    if hasher != expected_hash:
        messages.error(request, 'Link de confirmação inválido.')
        return redirect('sapl.materia:proposicao-confirmar', hash=hash, pk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = get_onlyoffice_browser_url(request)

    context = {
        'proposicao': proposicao,
        'hash': hash,
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.materia:onlyoffice_confirmar_config', kwargs={'pk': pk}),
        'voltar_url': reverse('sapl.materia:proposicao-confirmar', kwargs={'hash': hash, 'pk': pk}),
    }

    return render(request, 'materia/onlyoffice_confirmar_editor.html', context)
