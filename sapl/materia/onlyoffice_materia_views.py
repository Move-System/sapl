"""
Views e utilitários para integração com OnlyOffice Document Server
para Matéria Legislativa e Documento Acessório
"""
import hashlib
import json
import logging
import time
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.contrib import messages

from sapl.materia.models import MateriaLegislativa, DocumentoAcessorio
from sapl.utils import build_onlyoffice_url, get_onlyoffice_browser_url

logger = logging.getLogger(__name__)


def generate_file_key(prefix, doc_id, user_id):
    """
    Gera uma chave única para o documento no OnlyOffice
    A chave muda a cada edição para forçar o OnlyOffice a recarregar
    """
    timestamp = str(int(time.time()))
    string_to_hash = f"{prefix}_{doc_id}_user_{user_id}_{timestamp}"
    return hashlib.md5(string_to_hash.encode()).hexdigest()


# ============================================================
# Views para Matéria Legislativa
# ============================================================

@login_required
@require_http_methods(["GET"])
def materia_onlyoffice_config(request, pk):
    """
    Retorna a configuração JSON para inicializar o editor OnlyOffice
    """
    materia = get_object_or_404(MateriaLegislativa, pk=pk)

    # Verifica permissão de edição
    can_edit = request.user.has_perm('materia.change_materialegislativa')

    # URLs para o OnlyOffice acessar (dentro da rede Docker)
    download_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:materia_onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:materia_onlyoffice_callback', kwargs={'pk': pk})
    )


    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key("materia", materia.pk, request.user.pk),
        "title": f"Materia_{materia.pk}.docx",
        "url": download_url,
    }

    # Configuração do editor
    editor_config = {
        "mode": "edit" if can_edit else "view",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@require_http_methods(["GET"])
def materia_onlyoffice_download(request, pk):
    """
    Endpoint para o OnlyOffice baixar o documento
    Se não existe arquivo, retorna um documento em branco
    """
    materia = get_object_or_404(MateriaLegislativa, pk=pk)

    # Se já tem arquivo, retorna ele
    if materia.texto_original:
        try:
            with open(materia.texto_original.path, 'rb') as f:
                content = f.read()
            response = HttpResponse(content, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename="Materia_{pk}.docx"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo: {e}")

    # Tenta criar documento usando template
    try:
        from sapl.utils_template import criar_documento_com_template, criar_documento_em_branco

        # Busca tipo específico da matéria (TipoMateriaLegislativa)
        tipo_especifico = materia.tipo if materia.tipo else None

        file_stream = criar_documento_com_template(
            'materia',
            tipo_especifico,
            {
                'titulo': f'{materia.tipo} {materia.numero}/{materia.ano}',
                'descricao': materia.ementa,
                'tipo_display': 'Matéria Legislativa'
            }
        )

        # Se não encontrou template, cria documento em branco
        if not file_stream:
            file_stream = criar_documento_em_branco(
                f'{materia.tipo} {materia.numero}/{materia.ano}',
                materia.ementa,
                'Matéria Legislativa'
            )

        if file_stream:
            response = HttpResponse(
                file_stream.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="Materia_{pk}.docx"'
            return response

    except Exception as e:
        logger.error(f"Erro ao criar documento: {e}")

    # Fallback: Se python-docx não está instalado ou houve erro
    try:
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_heading(f'{materia.tipo} {materia.numero}/{materia.ano}', 0)
        doc.add_paragraph(f'Ementa: {materia.ementa}')
        doc.add_paragraph('')
        doc.add_paragraph('Digite o texto da matéria abaixo:')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="Materia_{pk}.docx"'
        return response

    except ImportError:
        logger.error("python-docx não está instalado")
        return HttpResponse("Erro: python-docx não instalado", status=500)


@csrf_exempt
@require_http_methods(["POST"])
def materia_onlyoffice_callback(request, pk):
    """
    Callback chamado pelo OnlyOffice quando o documento é salvo
    """
    try:
        body = json.loads(request.body.decode('utf-8'))
        status = body.get('status')
        download_url = body.get('url')

        logger.info(f"OnlyOffice callback para matéria legislativa {pk}: status={status}, url={download_url}")

        # Status 2 ou 6 significa que o documento foi salvo
        if status in [2, 6] and download_url:
            logger.info(f"URL original recebida: {download_url}")

            # Substitui o host da URL pelo ONLYOFFICE_URL configurado
            # O OnlyOffice retorna URLs com seu próprio hostname que pode não ser acessível
            import re
            from urllib.parse import urlparse
            onlyoffice_parsed = urlparse(settings.ONLYOFFICE_URL)
            onlyoffice_base = f"{onlyoffice_parsed.scheme}://{onlyoffice_parsed.netloc}"
            download_url = re.sub(
                r'https?://[^/]+',
                onlyoffice_base,
                download_url
            )
            logger.info(f"URL substituída para: {download_url}")

            materia = get_object_or_404(MateriaLegislativa, pk=pk)

            # Baixa o documento do OnlyOffice
            import requests
            try:
                response = requests.get(download_url, timeout=30)
                logger.info(f"Download response: status={response.status_code}, size={len(response.content)}")
            except Exception as e:
                logger.error(f"Erro ao fazer requisição de download: {e}")
                return JsonResponse({"error": 1})

            if response.status_code == 200:
                from django.core.files.base import ContentFile

                filename = f"materia_{pk}_{int(time.time())}.docx"

                # Remove arquivo antigo se existir
                if materia.texto_original:
                    old_file = materia.texto_original.name
                    materia.texto_original.delete(save=False)
                    logger.info(f"Arquivo antigo removido: {old_file}")

                try:
                    materia.texto_original.save(
                        filename,
                        ContentFile(response.content),
                        save=True
                    )
                    logger.info(f"Documento salvo com sucesso: {filename}")
                    return JsonResponse({"error": 0})
                except Exception as e:
                    logger.error(f"Erro ao salvar arquivo: {e}")
                    return JsonResponse({"error": 1})
            else:
                logger.error(f"Erro ao baixar documento: status={response.status_code}")
                return JsonResponse({"error": 1})

        return JsonResponse({"error": 0})

    except Exception as e:
        logger.error(f"Erro no callback OnlyOffice: {e}")
        return JsonResponse({"error": 1})


@login_required
@require_http_methods(["GET"])
def materia_check_doc(request, pk):
    """
    Verifica se a matéria já possui documento salvo.
    Usado pelo frontend para polling após forcesave do OnlyOffice.
    """
    materia = get_object_or_404(MateriaLegislativa, pk=pk)
    has_document = bool(materia.texto_original)
    return JsonResponse({"has_document": has_document})


@login_required
@require_http_methods(["POST"])
def materia_forcesave(request, pk):
    """
    Força o salvamento do documento via OnlyOffice Command Service.
    """
    get_object_or_404(MateriaLegislativa, pk=pk)

    body = json.loads(request.body.decode('utf-8'))
    doc_key = body.get('key')

    if not doc_key:
        return JsonResponse({"error": "key é obrigatório"}, status=400)

    command_url = f"{settings.ONLYOFFICE_URL}/coauthoring/CommandService.ashx"

    payload = {
        "c": "forcesave",
        "key": doc_key
    }

    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(payload, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        payload['token'] = token

    try:
        import requests as http_requests
        logger.info(f"Forcesave para matéria {pk}: key={doc_key}")
        resp = http_requests.post(command_url, json=payload, timeout=10)
        result = resp.json()
        logger.info(f"Forcesave response para matéria {pk}: {result}")
        return JsonResponse(result)
    except Exception as e:
        logger.error(f"Erro ao chamar forcesave para matéria {pk}: {e}")
        return JsonResponse({"error": str(e)}, status=500)


@login_required
def materia_onlyoffice_editor(request, pk):
    """
    Renderiza a página com o editor OnlyOffice integrado
    """
    materia = get_object_or_404(MateriaLegislativa, pk=pk)

    # Verifica se o usuário tem permissão
    # (administrador, operador de matéria, protocolo)
    if not (request.user.is_superuser or
            request.user.has_perm('materia.change_materialegislativa')):
        messages.error(request, 'Você não tem permissão para editar esta matéria.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = get_onlyoffice_browser_url(request)

    context = {
        'documento': materia,
        'documento_tipo': 'Matéria Legislativa',
        'documento_titulo': f'{materia.tipo} {materia.numero}/{materia.ano}',
        'documento_descricao': materia.ementa,
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.materia:materia_onlyoffice_config', kwargs={'pk': pk}),
        'voltar_url': reverse('sapl.materia:materialegislativa_detail', kwargs={'pk': pk}),
        'check_url': reverse('sapl.materia:materia_check_doc', kwargs={'pk': pk}),
        'forcesave_url': reverse('sapl.materia:materia_forcesave', kwargs={'pk': pk}),
    }

    return render(request, 'onlyoffice/onlyoffice_editor.html', context)


# ============================================================
# Views para Documento Acessório
# ============================================================

@login_required
@require_http_methods(["GET"])
def docacessorio_onlyoffice_config(request, pk):
    """
    Retorna a configuração JSON para inicializar o editor OnlyOffice
    """
    documento = get_object_or_404(DocumentoAcessorio, pk=pk)

    # Verifica permissão de edição
    can_edit = request.user.has_perm('materia.change_documentoacessorio')

    # URLs para o OnlyOffice acessar (dentro da rede Docker)
    download_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:docacessorio_onlyoffice_download', kwargs={'pk': pk})
    )
    callback_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:docacessorio_onlyoffice_callback', kwargs={'pk': pk})
    )


    # Configuração do documento
    document_config = {
        "fileType": "docx",
        "key": generate_file_key("docacessorio", documento.pk, request.user.pk),
        "title": f"DocAcessorio_{documento.pk}.docx",
        "url": download_url,
    }

    # Configuração do editor
    editor_config = {
        "mode": "edit" if can_edit else "view",
        "lang": "pt-BR",
        "callbackUrl": callback_url,
        "user": {
            "id": str(request.user.pk),
            "name": request.user.get_full_name() or request.user.username,
        },
        "customization": {
            "autosave": True,
            "forcesave": True,
            "comments": True,
            "chat": False,
        },
    }

    config = {
        "documentType": "word",
        "document": document_config,
        "editorConfig": editor_config,
        "height": "600px",
        "width": "100%",
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(config, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        config['token'] = token

    return JsonResponse(config)


@require_http_methods(["GET"])
def docacessorio_onlyoffice_download(request, pk):
    """
    Endpoint para o OnlyOffice baixar o documento
    Se não existe arquivo, retorna um documento em branco
    """
    documento = get_object_or_404(DocumentoAcessorio, pk=pk)

    # Se já tem arquivo, retorna ele
    if documento.arquivo:
        try:
            with open(documento.arquivo.path, 'rb') as f:
                content = f.read()
            response = HttpResponse(content, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename="DocAcessorio_{pk}.docx"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo: {e}")

    # Tenta criar documento usando template
    try:
        from sapl.utils_template import criar_documento_com_template, criar_documento_em_branco

        # Busca tipo específico do documento acessório (TipoDocumento)
        tipo_especifico = documento.tipo if documento.tipo else None

        descricao = documento.ementa if documento.ementa else f'Matéria: {documento.materia}'

        file_stream = criar_documento_com_template(
            'docacessorio',
            tipo_especifico,
            {
                'titulo': f'{documento.tipo} - {documento.nome}',
                'descricao': descricao,
                'tipo_display': 'Documento Acessório'
            }
        )

        # Se não encontrou template, cria documento em branco
        if not file_stream:
            file_stream = criar_documento_em_branco(
                f'{documento.tipo} - {documento.nome}',
                descricao,
                'Documento Acessório'
            )

        if file_stream:
            response = HttpResponse(
                file_stream.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename="DocAcessorio_{pk}.docx"'
            return response

    except Exception as e:
        logger.error(f"Erro ao criar documento: {e}")

    # Fallback: Se python-docx não está instalado ou houve erro
    try:
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_heading(f'{documento.tipo} - {documento.nome}', 0)
        if documento.ementa:
            doc.add_paragraph(f'Ementa: {documento.ementa}')
        doc.add_paragraph(f'Matéria: {documento.materia}')
        doc.add_paragraph('')
        doc.add_paragraph('Digite o texto do documento abaixo:')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="DocAcessorio_{pk}.docx"'
        return response

    except ImportError:
        logger.error("python-docx não está instalado")
        return HttpResponse("Erro: python-docx não instalado", status=500)


@csrf_exempt
@require_http_methods(["POST"])
def docacessorio_onlyoffice_callback(request, pk):
    """
    Callback chamado pelo OnlyOffice quando o documento é salvo
    """
    try:
        body = json.loads(request.body.decode('utf-8'))
        status = body.get('status')
        download_url = body.get('url')

        logger.info(f"OnlyOffice callback para documento acessório {pk}: status={status}, url={download_url}")

        # Status 2 ou 6 significa que o documento foi salvo
        if status in [2, 6] and download_url:
            logger.info(f"URL original recebida: {download_url}")

            # Substitui o host da URL pelo ONLYOFFICE_URL configurado
            # O OnlyOffice retorna URLs com seu próprio hostname que pode não ser acessível
            import re
            from urllib.parse import urlparse
            onlyoffice_parsed = urlparse(settings.ONLYOFFICE_URL)
            onlyoffice_base = f"{onlyoffice_parsed.scheme}://{onlyoffice_parsed.netloc}"
            download_url = re.sub(
                r'https?://[^/]+',
                onlyoffice_base,
                download_url
            )
            logger.info(f"URL substituída para: {download_url}")

            documento = get_object_or_404(DocumentoAcessorio, pk=pk)

            # Baixa o documento do OnlyOffice
            import requests
            try:
                response = requests.get(download_url, timeout=30)
                logger.info(f"Download response: status={response.status_code}, size={len(response.content)}")
            except Exception as e:
                logger.error(f"Erro ao fazer requisição de download: {e}")
                return JsonResponse({"error": 1})

            if response.status_code == 200:
                from django.core.files.base import ContentFile

                filename = f"docacessorio_{pk}_{int(time.time())}.docx"

                # Remove arquivo antigo se existir
                if documento.arquivo:
                    old_file = documento.arquivo.name
                    documento.arquivo.delete(save=False)
                    logger.info(f"Arquivo antigo removido: {old_file}")

                try:
                    documento.arquivo.save(
                        filename,
                        ContentFile(response.content),
                        save=True
                    )
                    logger.info(f"Documento salvo com sucesso: {filename}")
                    return JsonResponse({"error": 0})
                except Exception as e:
                    logger.error(f"Erro ao salvar arquivo: {e}")
                    return JsonResponse({"error": 1})
            else:
                logger.error(f"Erro ao baixar documento: status={response.status_code}")
                return JsonResponse({"error": 1})

        return JsonResponse({"error": 0})

    except Exception as e:
        logger.error(f"Erro no callback OnlyOffice: {e}")
        return JsonResponse({"error": 1})


@login_required
@require_http_methods(["GET"])
def docacessorio_check_doc(request, pk):
    """
    Verifica se o documento acessório já possui arquivo salvo.
    Usado pelo frontend para polling após forcesave do OnlyOffice.
    """
    documento = get_object_or_404(DocumentoAcessorio, pk=pk)
    has_document = bool(documento.arquivo)
    return JsonResponse({"has_document": has_document})


@login_required
@require_http_methods(["POST"])
def docacessorio_forcesave(request, pk):
    """
    Força o salvamento do documento via OnlyOffice Command Service.
    """
    get_object_or_404(DocumentoAcessorio, pk=pk)

    body = json.loads(request.body.decode('utf-8'))
    doc_key = body.get('key')

    if not doc_key:
        return JsonResponse({"error": "key é obrigatório"}, status=400)

    command_url = f"{settings.ONLYOFFICE_URL}/coauthoring/CommandService.ashx"

    payload = {
        "c": "forcesave",
        "key": doc_key
    }

    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(payload, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        payload['token'] = token

    try:
        import requests as http_requests
        logger.info(f"Forcesave para documento acessório {pk}: key={doc_key}")
        resp = http_requests.post(command_url, json=payload, timeout=10)
        result = resp.json()
        logger.info(f"Forcesave response para documento acessório {pk}: {result}")
        return JsonResponse(result)
    except Exception as e:
        logger.error(f"Erro ao chamar forcesave para documento acessório {pk}: {e}")
        return JsonResponse({"error": str(e)}, status=500)


@login_required
def docacessorio_onlyoffice_editor(request, pk):
    """
    Renderiza a página com o editor OnlyOffice integrado
    """
    documento = get_object_or_404(DocumentoAcessorio, pk=pk)

    # Verifica se o usuário tem permissão
    # (administrador, operador de matéria, protocolo)
    if not (request.user.is_superuser or
            request.user.has_perm('materia.change_documentoacessorio')):
        messages.error(request, 'Você não tem permissão para editar este documento.')
        return redirect('sapl.materia:documentoacessorio_detail', pk=documento.materia.pk, zpk=pk)

    # URL do OnlyOffice acessível pelo navegador do usuário
    onlyoffice_url = get_onlyoffice_browser_url(request)

    context = {
        'documento': documento,
        'documento_tipo': 'Documento Acessório',
        'documento_titulo': f'{documento.tipo} - {documento.nome}',
        'documento_descricao': documento.ementa or f'Matéria: {documento.materia}',
        'onlyoffice_url': onlyoffice_url,
        'config_url': reverse('sapl.materia:docacessorio_onlyoffice_config', kwargs={'pk': pk}),
        'voltar_url': reverse('sapl.materia:documentoacessorio_detail', kwargs={'pk': documento.materia.pk, 'zpk': pk}),
        'check_url': reverse('sapl.materia:docacessorio_check_doc', kwargs={'pk': pk}),
        'forcesave_url': reverse('sapl.materia:docacessorio_forcesave', kwargs={'pk': pk}),
    }

    return render(request, 'onlyoffice/onlyoffice_editor.html', context)


# ============================================================
# Views para Geração de PDF para Assinatura
# ============================================================

@login_required
@require_http_methods(["GET"])
def materia_gerar_pdf_assinatura(request, pk):
    """
    Gera o PDF do documento da matéria para assinatura digital.
    Usa a API de conversão do OnlyOffice para converter DOCX para PDF.
    """
    import requests as http_requests
    import xml.etree.ElementTree as ET

    materia = get_object_or_404(MateriaLegislativa, pk=pk)

    # Verifica se a matéria tem número de protocolo
    if not materia.numero_protocolo:
        messages.error(request, 'Esta matéria ainda não possui número de protocolo.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)

    # Verifica se existe documento
    if not materia.texto_original:
        messages.error(request, 'Esta matéria não possui documento de texto original.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)

    # Verifica a extensão do arquivo
    file_path = materia.texto_original.path
    file_name = materia.texto_original.name.lower()

    # Se já é PDF, retorna diretamente
    if file_name.endswith('.pdf'):
        try:
            with open(file_path, 'rb') as f:
                content = f.read()
            response = HttpResponse(content, content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="Materia_{materia.tipo}_{materia.numero}_{materia.ano}.pdf"'
            return response
        except Exception as e:
            logger.error(f"Erro ao ler arquivo PDF: {e}")
            messages.error(request, 'Erro ao ler o arquivo PDF.')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

    # URL do documento para o OnlyOffice (dentro da rede Docker)
    download_url = build_onlyoffice_url(
        request,
        reverse('sapl.materia:materia_onlyoffice_download', kwargs={'pk': pk})
    )

    # URL da API de conversão do OnlyOffice
    conversion_url = f'{settings.ONLYOFFICE_URL}/ConvertService.ashx'

    # Configuração da conversão
    conversion_data = {
        "async": False,
        "filetype": "docx",
        "key": generate_file_key("materia_pdf", materia.pk, request.user.pk),
        "outputtype": "pdf",
        "title": f"Materia_{materia.tipo}_{materia.numero}_{materia.ano}.pdf",
        "url": download_url,
    }

    # Adiciona JWT se estiver habilitado
    if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
        import jwt
        token = jwt.encode(conversion_data, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        conversion_data['token'] = token

    try:
        # Solicita a conversão
        headers = {'Content-Type': 'application/json'}

        # Adiciona header de autorização JWT se habilitado
        if settings.ONLYOFFICE_JWT_ENABLED and settings.ONLYOFFICE_JWT_SECRET:
            import jwt
            header_token = jwt.encode({"payload": conversion_data}, settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
            headers['Authorization'] = f'Bearer {header_token}'

        conversion_response = http_requests.post(
            conversion_url,
            json=conversion_data,
            headers=headers,
            timeout=60
        )

        if conversion_response.status_code != 200:
            logger.error(f"Erro na conversão OnlyOffice: status={conversion_response.status_code}")
            messages.error(request, 'Erro ao converter o documento para PDF.')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

        # A API do OnlyOffice retorna XML
        try:
            root = ET.fromstring(conversion_response.text)
        except ET.ParseError as e:
            logger.error(f"Erro ao parsear resposta XML do OnlyOffice: {e}")
            logger.error(f"Resposta recebida: {conversion_response.text[:500]}")
            messages.error(request, 'Erro ao processar resposta do serviço de conversão.')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

        # Verifica se houve erro
        error_elem = root.find('Error')
        if error_elem is not None:
            error_code = error_elem.text
            logger.error(f"Erro na conversão OnlyOffice: error={error_code}")
            messages.error(request, f'Erro na conversão do documento: código {error_code}')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

        # Obtém a URL do PDF convertido
        file_url_elem = root.find('FileUrl')
        if file_url_elem is None or not file_url_elem.text:
            logger.error("URL do PDF não retornada pelo OnlyOffice")
            logger.error(f"Resposta XML: {conversion_response.text}")
            messages.error(request, 'Erro ao obter o documento PDF convertido.')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

        pdf_url = file_url_elem.text
        logger.info(f"URL do PDF para download: {pdf_url}")

        # Baixa o PDF convertido
        pdf_response = http_requests.get(pdf_url, timeout=60)

        if pdf_response.status_code != 200:
            logger.error(f"Erro ao baixar PDF convertido: status={pdf_response.status_code}")
            messages.error(request, 'Erro ao baixar o documento PDF convertido.')
            return redirect('sapl.materia:materialegislativa_detail', pk=pk)

        # Retorna o PDF
        filename = f"Materia_{materia.tipo}_{materia.numero}_{materia.ano}.pdf"
        # Remove caracteres especiais do nome do arquivo
        filename = filename.replace(' ', '_').replace('/', '-')

        response = HttpResponse(pdf_response.content, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response

    except http_requests.exceptions.Timeout:
        logger.error("Timeout na conversão OnlyOffice")
        messages.error(request, 'Tempo limite excedido ao converter o documento.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)
    except http_requests.exceptions.ConnectionError:
        logger.error("Erro de conexão com OnlyOffice")
        messages.error(request, 'Não foi possível conectar ao serviço de conversão de documentos.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)
    except Exception as e:
        logger.error(f"Erro inesperado na geração de PDF: {e}")
        messages.error(request, 'Erro inesperado ao gerar o PDF.')
        return redirect('sapl.materia:materialegislativa_detail', pk=pk)
