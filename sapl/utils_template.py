"""
Utilitários para aplicar templates de documentos aos novos documentos
criados via OnlyOffice.
"""
import logging
import os
from io import BytesIO

from django.core.files.base import ContentFile

logger = logging.getLogger(__name__)


def _forcar_fonte_estilo(style):
    """Força Arial em todos os slots de fonte de um estilo,
    sobrescrevendo qualquer fonte de tema."""
    from docx.oxml.ns import qn

    style.font.name = 'Arial'
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rFonts.set(qn('w:eastAsia'), 'Arial')
    rFonts.set(qn('w:cs'), 'Arial')


def criar_documento_padrao():
    """
    Cria um Document python-docx com fonte Arial como padrão.
    Configura o estilo 'Normal', 'Title' e os estilos de heading para usar Arial,
    forçando em todos os slots de fonte (ascii, hAnsi, eastAsia, cs).
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Define Arial como fonte padrão no estilo Normal
    style = doc.styles['Normal']
    style.font.size = Pt(12)
    _forcar_fonte_estilo(style)

    # Define Arial no estilo Title (usado por add_heading(text, 0))
    if 'Title' in doc.styles:
        _forcar_fonte_estilo(doc.styles['Title'])

    # Define Arial nos estilos de heading
    for i in range(1, 10):
        heading_style_name = f'Heading {i}'
        if heading_style_name in doc.styles:
            _forcar_fonte_estilo(doc.styles[heading_style_name])

    return doc


def criar_documento_com_template(tipo_conteudo, tipo_especifico, dados):
    """
    Cria um novo documento usando um template existente.

    1. Busca template adequado via DocumentTemplate.get_template_for()
    2. Carrega .docx do template com python-docx
    3. Mantém cabeçalho/rodapé, limpa ou adapta o corpo
    4. Adiciona dados iniciais do documento
    5. Retorna BytesIO com documento pronto

    Args:
        tipo_conteudo: string com o tipo (proposicao, materia, docacessorio, docadm, norma)
        tipo_especifico: objeto do tipo específico (TipoMateriaLegislativa, TipoProposicao, etc.)
                        ou None para usar template genérico
        dados: dicionário com dados do documento:
            - titulo: título ou identificação do documento
            - descricao: descrição, ementa ou assunto
            - tipo_display: nome do tipo para exibição (opcional)

    Returns:
        BytesIO com o documento pronto ou None se não houver template
    """
    try:
        from docx import Document
        from docx.shared import Pt
        from sapl.base.models import DocumentTemplate

        # Busca o template mais adequado
        template = DocumentTemplate.get_template_for(tipo_conteudo, tipo_especifico)

        if not template or not template.arquivo:
            logger.debug(f"Nenhum template encontrado para {tipo_conteudo}")
            return None

        logger.info(f"Usando template '{template.nome}' para {tipo_conteudo}")

        # Carrega o documento do template
        try:
            doc = Document(template.arquivo.path)
        except Exception as e:
            logger.error(f"Erro ao carregar template {template.pk}: {e}")
            return None

        # Substitui placeholders no corpo do documento
        # Placeholders suportados: {titulo}, {ementa}, {descricao}, {tipo}, {data}
        titulo = dados.get('titulo', 'Documento')
        descricao = dados.get('descricao', '')
        tipo_display = dados.get('tipo_display', tipo_conteudo.title())

        from datetime import date
        data_atual = date.today().strftime('%d/%m/%Y')

        placeholders = {
            '{titulo}': titulo,
            '{ementa}': descricao,
            '{descricao}': descricao,
            '{tipo}': tipo_display,
            '{data}': data_atual,
        }

        # Substitui placeholders nos parágrafos
        for paragraph in doc.paragraphs:
            for placeholder, valor in placeholders.items():
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, valor)

        # Substitui placeholders nas tabelas
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder, valor in placeholders.items():
                            if placeholder in paragraph.text:
                                paragraph.text = paragraph.text.replace(placeholder, valor)

        # Substitui placeholders no cabeçalho e rodapé
        for section in doc.sections:
            for header_para in section.header.paragraphs:
                for placeholder, valor in placeholders.items():
                    if placeholder in header_para.text:
                        header_para.text = header_para.text.replace(placeholder, valor)
            for footer_para in section.footer.paragraphs:
                for placeholder, valor in placeholders.items():
                    if placeholder in footer_para.text:
                        footer_para.text = footer_para.text.replace(placeholder, valor)

        # Salva em memória
        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return file_stream

    except ImportError:
        logger.error("python-docx não está instalado")
        return None
    except Exception as e:
        logger.error(f"Erro ao criar documento com template: {e}")
        return None


def criar_documento_em_branco(titulo, descricao, tipo_documento='Documento'):
    """
    Cria um documento em branco básico quando não há template disponível.
    Fallback para manter compatibilidade com o comportamento anterior.

    Args:
        titulo: título ou identificação do documento
        descricao: descrição, ementa ou assunto
        tipo_documento: tipo para exibição (ex: 'Proposição', 'Matéria Legislativa')

    Returns:
        BytesIO com o documento ou None em caso de erro
    """
    try:
        from io import BytesIO

        doc = criar_documento_padrao()
        doc.add_heading(titulo, 0)

        if descricao:
            doc.add_paragraph(f'Ementa: {descricao}')

        doc.add_paragraph('')
        doc.add_paragraph(f'Digite o texto do documento abaixo:')
        doc.add_paragraph('')

        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return file_stream

    except ImportError:
        logger.error("python-docx não está instalado")
        return None
    except Exception as e:
        logger.error(f"Erro ao criar documento em branco: {e}")
        return None


def aplicar_template_existente(documento_path, tipo_conteudo, tipo_especifico=None):
    """
    Aplica cabeçalho e rodapé de um template a um documento existente.

    Útil para documentos que já foram criados mas precisam ter o
    cabeçalho/rodapé atualizado.

    Args:
        documento_path: caminho para o documento .docx existente
        tipo_conteudo: string com o tipo (proposicao, materia, etc.)
        tipo_especifico: objeto do tipo específico ou None

    Returns:
        BytesIO com o documento atualizado ou None se não houver template
    """
    try:
        from docx import Document
        from sapl.base.models import DocumentTemplate
        from copy import deepcopy

        # Busca o template
        template = DocumentTemplate.get_template_for(tipo_conteudo, tipo_especifico)

        if not template or not template.arquivo:
            logger.debug(f"Nenhum template encontrado para aplicar")
            return None

        # Carrega o template e o documento existente
        template_doc = Document(template.arquivo.path)
        doc = Document(documento_path)

        # Copia cabeçalho do template para o documento
        if template_doc.sections:
            template_section = template_doc.sections[0]

            for section in doc.sections:
                # Copia cabeçalho
                if template_section.header:
                    for para in section.header.paragraphs:
                        p = para._element
                        p.getparent().remove(p)

                    for para in template_section.header.paragraphs:
                        new_para = section.header.add_paragraph()
                        new_para.text = para.text
                        new_para.style = para.style
                        new_para.alignment = para.alignment

                # Copia rodapé
                if template_section.footer:
                    for para in section.footer.paragraphs:
                        p = para._element
                        p.getparent().remove(p)

                    for para in template_section.footer.paragraphs:
                        new_para = section.footer.add_paragraph()
                        new_para.text = para.text
                        new_para.style = para.style
                        new_para.alignment = para.alignment

        # Salva em memória
        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return file_stream

    except ImportError:
        logger.error("python-docx não está instalado")
        return None
    except Exception as e:
        logger.error(f"Erro ao aplicar template a documento existente: {e}")
        return None


def adicionar_cabecalho_materia(materia):
    """
    Adiciona cabeçalho com identificação da matéria no início do documento.

    Adiciona um parágrafo no início do documento com o formato:
    "TIPO NÚMERO / ANO" (ex: "INDICAÇÃO 10 / 2026")

    Args:
        materia: objeto MateriaLegislativa com texto_original preenchido

    Returns:
        True se o cabeçalho foi adicionado com sucesso, False caso contrário
    """
    if not materia.texto_original:
        logger.warning(f"Matéria {materia.pk} não possui texto_original")
        return False

    # Verifica se é um arquivo .docx
    arquivo_path = materia.texto_original.path
    extensao = os.path.splitext(arquivo_path)[1].lower()

    if extensao not in ['.docx', '.doc']:
        logger.info(f"Arquivo {arquivo_path} não é .docx, pulando adição de cabeçalho")
        return False

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        def _forcar_fonte_arial(run):
            """Força Arial em todos os slots de fonte do run,
            sobrescrevendo qualquer fonte de tema do documento."""
            run.bold = True
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:ascii'), 'Arial')
            rFonts.set(qn('w:hAnsi'), 'Arial')
            rFonts.set(qn('w:eastAsia'), 'Arial')
            rFonts.set(qn('w:cs'), 'Arial')

        # Carrega o documento
        doc = Document(arquivo_path)

        # Monta o texto do cabeçalho: "TIPO NÚMERO / ANO"
        tipo_nome = str(materia.tipo).upper()
        cabecalho_texto = f"{tipo_nome} {materia.numero} / {materia.ano}"

        # Insere o cabeçalho no início do documento
        # Precisamos inserir antes do primeiro parágrafo
        if doc.paragraphs:
            primeiro_paragrafo = doc.paragraphs[0]

            # Cria novo parágrafo antes do primeiro
            novo_paragrafo = primeiro_paragrafo.insert_paragraph_before(cabecalho_texto)

            # Formata o cabeçalho
            novo_paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in novo_paragrafo.runs:
                _forcar_fonte_arial(run)

            # Adiciona linha em branco após o cabeçalho
            primeiro_paragrafo.insert_paragraph_before('')
        else:
            # Documento vazio, adiciona o cabeçalho como primeiro parágrafo
            paragrafo = doc.add_paragraph(cabecalho_texto)
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragrafo.runs:
                _forcar_fonte_arial(run)
            doc.add_paragraph('')

        # Salva o documento modificado
        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        # Atualiza o arquivo da matéria
        nome_arquivo = os.path.basename(arquivo_path)
        materia.texto_original.save(nome_arquivo, ContentFile(file_stream.read()), save=True)

        logger.info(f"Cabeçalho adicionado à matéria {materia.pk}: {cabecalho_texto}")
        return True

    except ImportError:
        logger.error("python-docx não está instalado")
        return False
    except Exception as e:
        logger.error(f"Erro ao adicionar cabeçalho à matéria {materia.pk}: {e}")
        return False
